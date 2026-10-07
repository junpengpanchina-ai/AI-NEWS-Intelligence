import csv
import io
import json
import logging
import re
import time
import uuid
from datetime import datetime, timezone
from urllib.parse import urlparse
import ipaddress

import httpx

from app.db import connect
from app.google_search import google_cse_card, provider_health, serp_choices
from app.ledger import create_source, get_source
from app.trace import record_trace

_log = logging.getLogger("app.intake")

CRAWL_TYPES = {
    "landing_page",
    "pricing_page",
    "docs_page",
    "blog_page",
    "sitemap",
    "rss",
    "checkout_signal",
}
DATASET_TYPES = {
    "stripe_payment_ranking": ("payment_signal", "Stripe"),
    "payment_ranking": ("payment_signal", "manual_csv"),
    "dr_growth_ranking": ("authority_signal", "sitedata"),
    "traffic_growth_ranking": ("traffic_signal", "sitedata"),
    "new_website_ranking": ("market_signal", ""),
    "serp_result": ("serp_result", "serper"),
    "keyword_signal": ("keyword_signal", ""),
    "google_trends": ("trend_signal", "google_trends"),
    "manual_intelligence": ("external_news", "manual"),
}
_DATASET_SOURCE = {
    "stripe_payment_ranking": ("Stripe Payment Ranking", "payment_ranking", "Stripe"),
    "payment_ranking": ("Payment Ranking", "payment_ranking", "manual_csv"),
    "dr_growth_ranking": ("DR Growth", "dr_growth", "sitedata"),
    "traffic_growth_ranking": ("SiteData Traffic Growth", "site_traffic", "SiteData"),
    "new_website_ranking": ("SiteData New Site Growth", "new_site_growth", "SiteData"),
    "serp_result": ("Manual Research", "serp", "serper"),
    "keyword_signal": ("Manual Research", "manual", "Manual"),
    "google_trends": ("Google Trends", "google_trends", "google_trends"),
    "manual_intelligence": ("Manual Intelligence", "manual", "manual"),
    "public_web_crawl": ("Public Web Crawl", "public_web", "Public Web"),
}
_PAYMENT_WORDS = ("stripe", "paddle", "dodo", "lemonsqueezy", "paypal", "checkout", "billing", "subscribe", "pricing")
_MONTH_FILE = re.compile(r"(20\d{2})-(\d{2})")
_MONTH_CN = re.compile(r"(20\d{2})年(\d{1,2})月")
_MAX_FILES = 40
_MAX_BYTES = 8_000_000
_MAX_ROWS = 5000
_SITEMAP_CAP = 20
_RECORD_TYPES = {
    "payment_signal",
    "authority_signal",
    "traffic_signal",
    "market_signal",
    "serp_result",
    "keyword_signal",
    "validation_signal",
    "crawl_signal",
    "trend_signal",
    "external_news",
}
_previews: dict[str, dict] = {}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _compact(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def _shift_month(period: str, delta: int) -> str:
    year = int(period[:4])
    month = int(period[5:7]) + delta
    while month < 1:
        month += 12
        year -= 1
    while month > 12:
        month -= 12
        year += 1
    return f"{year:04d}-{month:02d}"


def _months_in_text(text: str) -> list[str]:
    found = []
    for match in _MONTH_CN.finditer(_compact(text)):
        found.append(f"{int(match.group(1)):04d}-{int(match.group(2)):02d}")
    return found


def _period_from_name(name: str) -> str:
    match = _MONTH_FILE.search(name or "")
    if not match:
        return ""
    month = int(match.group(2))
    if month < 1 or month > 12:
        return ""
    return f"{match.group(1)}-{month:02d}"


def _header_keys(headers: list[str]) -> set[str]:
    return {(header or "").strip().lower().replace(" ", "_").replace("-", "_") for header in headers}


def detect_dataset(file_name: str, relative_path: str, headers: list[str]) -> dict:
    name = (file_name or "").lower()
    path = (relative_path or "").lower()
    blob = f"{path} {name}"
    warnings = []

    def pack(dataset: str, confidence: float) -> dict:
        record_type, provider = DATASET_TYPES.get(dataset, ("", ""))
        return {
            "detected_dataset_type": dataset,
            "record_type": record_type,
            "provider": provider,
            "confidence": confidence,
            "warnings": warnings,
        }

    if "stripe-payment-ranking" in name or _stripe_headers(headers):
        return pack("stripe_payment_ranking", 0.95 if "stripe-payment-ranking" in name else 0.86)
    if "dr-fastest-growing-websites" in name or _dr_headers(headers):
        return pack("dr_growth_ranking", 0.95 if "dr-fastest-growing-websites" in name else 0.86)
    if "新网站排名" in blob or "new-websites" in blob:
        return pack("new_website_ranking", 0.95)
    if ("fastest-growing-websites" in name and "dr" not in name) or _traffic_headers(headers):
        confidence = 0.95 if "fastest-growing-websites" in name and "dr" not in name else 0.86
        return pack("traffic_growth_ranking", confidence)
    keys = _header_keys(headers)
    if "keyword" in keys and keys & {"search_volume", "growth_rate", "related_queries", "trend_url", "started_at"}:
        return pack("google_trends", 0.93)
    if {"domain", "payment_traffic"} <= keys or "payment_provider" in keys:
        found = pack("payment_ranking", 0.9)
        if "dodo" in blob:
            found["provider"] = "dodo"
        elif "nexi" in blob:
            found["provider"] = "nexi"
        elif "stripe" in blob:
            found["provider"] = "stripe"
        return found
    if {"domain", "dr", "dr_growth"} <= keys:
        return pack("dr_growth_ranking", 0.9)
    if {"domain", "traffic"} <= keys and keys & {"growth_rate", "month", "category"}:
        return pack("traffic_growth_ranking", 0.88)
    if {"title", "url"} <= keys and "rank" not in keys and keys & {"source", "published_at", "note", "category"}:
        found = pack("manual_intelligence", 0.9)
        if "producthunt" in blob or "product_hunt" in blob:
            found["provider"] = "producthunt"
        elif "hacker" in blob or "/hn" in blob:
            found["provider"] = "hn"
        elif "github" in blob:
            found["provider"] = "github"
        elif "rss" in blob:
            found["provider"] = "rss"
        return found
    if {"title", "url", "domain", "snippet", "rank"} <= keys or {"query", "rank", "title", "url"} <= keys:
        found = pack("serp_result", 0.9)
        if "google" in blob and "cse" in blob:
            found["provider"] = "google_cse"
        elif "manual" in blob:
            found["provider"] = "manual_csv"
        return found
    if {"keyword", "score", "intent", "page_type"} <= keys:
        return pack("keyword_signal", 0.9)
    warnings.append("未能识别数据类型")
    return pack("unknown", 0.0)


def _stripe_headers(headers: list[str]) -> bool:
    for header in headers:
        lower = (header or "").lower()
        compact = _compact(header).lower()
        if "导向" in header and "stripe" in lower:
            return True
        if compact in {"domainrating", "domain_rating"} or lower.strip() == "domain rating":
            return True
        if "google ads advertiser sites" in lower:
            return True
    return False


def _dr_headers(headers: list[str]) -> bool:
    for header in headers:
        compact = _compact(header)
        if "DR增长" in compact:
            return True
        if "上月" in header and "域名评分" in header:
            return True
        if _MONTH_CN.search(compact) and compact.upper().endswith("DR"):
            return True
    return False


def _traffic_headers(headers: list[str]) -> bool:
    for header in headers:
        compact = _compact(header)
        if "流量增长" in compact or "增长率" in compact:
            return True
        if _MONTH_CN.search(compact) and compact.endswith("流量") and "上月" not in header:
            return True
    return False


def _read_csv(data: bytes) -> tuple[list[str], list[dict]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV 需要是 UTF-8 或 UTF-8-SIG") from exc
    if not text.strip():
        return [], []
    sample = text.splitlines()[0]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    header_row = next(reader, None)
    if header_row is None:
        return [], []
    headers = _unique_headers(header_row)
    return headers, _rows_from(headers, reader)


def _read_xlsx(data: bytes) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            return [], []
        headers = _unique_headers(["" if cell is None else str(cell) for cell in header_row])
        rows = []
        for cells in iterator:
            values = ["" if cell is None else str(cell).strip() for cell in cells]
            if not any(values):
                continue
            rows.append({header: values[index] if index < len(values) else "" for index, header in enumerate(headers)})
            if len(rows) >= _MAX_ROWS:
                break
        return headers, rows
    finally:
        workbook.close()


def _unique_headers(cells: list[str]) -> list[str]:
    headers = []
    seen = {}
    for index, cell in enumerate(cells):
        name = (cell or "").strip() or f"column_{index + 1}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        headers.append(name if count == 1 else f"{name}_{count}")
    return headers


def _rows_from(headers: list[str], reader) -> list[dict]:
    rows = []
    for cells in reader:
        if not any((cell or "").strip() for cell in cells):
            continue
        rows.append({header: (cells[index].strip() if index < len(cells) and cells[index] else "") for index, header in enumerate(headers)})
        if len(rows) >= _MAX_ROWS:
            break
    return rows


def _read_table(file_name: str, data: bytes) -> tuple[list[str], list[dict], list[str]]:
    warnings = []
    lower = (file_name or "").lower()
    if lower.endswith(".xlsx"):
        headers, rows = _read_xlsx(data)
    elif lower.endswith(".csv") or not lower.endswith(".xls"):
        headers, rows = _read_csv(data)
    else:
        raise ValueError("只支持 CSV 或 XLSX")
    if len(rows) >= _MAX_ROWS:
        warnings.append(f"仅保留前 {_MAX_ROWS} 行")
    return headers, rows, warnings


def _resolve_months(file_name: str, relative_path: str, headers: list[str]) -> tuple[str, str]:
    period = _period_from_name(file_name) or _period_from_name(relative_path)
    header_months = []
    for header in headers:
        header_months.extend(_months_in_text(header))
    unique = []
    for month in header_months:
        if month not in unique:
            unique.append(month)
    unique.sort()
    if not period and unique:
        period = unique[-1]
    previous = ""
    if period:
        earlier = [month for month in unique if month < period]
        previous = earlier[-1] if earlier else _shift_month(period, -1)
    return period, previous


def _map_header(dataset: str, header: str, period: str) -> str:
    text = header or ""
    lower = text.lower().strip()
    compact = _compact(text)
    compact_lower = compact.lower()
    month = ""
    match = _MONTH_CN.search(compact)
    if match:
        month = f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"
    if dataset == "stripe_payment_ranking":
        if lower == "rank":
            return "rank"
        if lower == "website":
            return "domain"
        if lower == "title":
            return "title"
        if "导向" in text and "stripe" in lower:
            return "payment_traffic"
        if month and compact.endswith("流量"):
            return "monthly_traffic"
        if "自然流量占比" in compact:
            return "organic_traffic_share"
        if compact_lower in {"domainrating", "domain_rating"} or lower == "domain rating":
            return "domain_rating"
        if "adsense" in lower:
            return "google_adsense_site_count"
        if "advertiser" in lower:
            return "google_ads_advertiser_site_count"
        if lower == "registration date":
            return "registered_at"
    if dataset == "dr_growth_ranking":
        if text == "排名":
            return "rank"
        if text == "网站":
            return "domain"
        if text == "标题":
            return "title"
        if month and compact.upper().endswith("DR"):
            return "current_dr" if not period or month == period else "previous_dr"
        if "上月" in text and "域名评分" in text:
            return "previous_dr"
        if "DR增长" in compact:
            return "dr_growth"
        if month and compact.endswith("流量"):
            return "current_traffic"
        if "adsense" in lower or "AdSense" in text:
            return "google_adsense_site_count"
        if "广告商" in text:
            return "google_ads_advertiser_site_count"
        if "注册日期" in text:
            return "registered_at"
        if lower in {"domain", "website"}:
            return "domain"
        if lower == "dr":
            return "current_dr"
        if lower == "dr_growth":
            return "dr_growth"
        if lower == "month":
            return "period"
        if lower == "category":
            return "category"
    if dataset == "traffic_growth_ranking":
        if text == "排名":
            return "rank"
        if text == "网站":
            return "domain"
        if text == "标题":
            return "title"
        if "上月" in text and "流量" in text:
            return "previous_traffic"
        if month and compact.endswith("流量"):
            return "current_traffic" if not period or month == period else "previous_traffic"
        if "流量增长" in compact:
            return "traffic_growth"
        if "增长率" in compact:
            return "growth_rate"
        if "域名评分" in text:
            return "domain_score"
        if "adsense" in lower or "AdSense" in text:
            return "google_adsense_site_count"
        if "广告商" in text:
            return "google_ads_advertiser_site_count"
        if "注册日期" in text:
            return "registered_at"
        if lower in {"domain", "website"}:
            return "domain"
        if lower == "traffic":
            return "traffic"
        if lower in {"growth_rate", "growth"}:
            return "growth_rate"
        if lower == "month":
            return "period"
        if lower == "category":
            return "category"
    if dataset == "payment_ranking":
        return {
            "domain": "domain",
            "website": "domain",
            "payment_traffic": "payment_traffic",
            "rank": "rank",
            "month": "period",
            "payment_provider": "payment_provider",
            "provider": "payment_provider",
            "category": "category",
        }.get(lower.replace(" ", "_"), "")
    if dataset == "google_trends":
        return {
            "keyword": "keyword",
            "search_volume": "search_volume",
            "started_at": "started_at",
            "growth_rate": "growth_rate",
            "growth": "growth_rate",
            "region": "region",
            "related_queries": "related_queries",
            "trend_url": "trend_url",
            "period": "period",
            "month": "period",
        }.get(lower.replace(" ", "_"), "")
    if dataset == "manual_intelligence":
        return {
            "title": "title",
            "url": "url",
            "domain": "domain",
            "source": "source_name",
            "published_at": "published_at",
            "note": "note",
            "category": "category",
        }.get(lower.replace(" ", "_"), "")
    if dataset == "new_website_ranking":
        if text in {"排名", "Rank"} or lower == "rank":
            return "rank"
        if text in {"网站", "Website"} or lower in {"website", "domain"}:
            return "domain"
        if text in {"标题", "Title"} or lower == "title":
            return "title"
        if text == "流量" or lower == "traffic":
            return "traffic"
        if "域名评分" in text or lower in {"domain rating", "domain_score"}:
            return "domain_score"
        if "注册日期" in text or lower == "registration date":
            return "registered_at"
    if dataset == "serp_result":
        return {
            "query": "keyword",
            "title": "title",
            "url": "url",
            "domain": "domain",
            "snippet": "snippet",
            "rank": "rank",
            "position": "rank",
        }.get(lower.replace(" ", "_"), "")
    if dataset == "keyword_signal":
        return {
            "keyword": "keyword",
            "score": "score",
            "intent": "intent",
            "page_type": "page_type",
        }.get(lower.replace(" ", "_"), "")
    return ""


def _map_row(dataset: str, row: dict, meta: dict) -> dict:
    mapped = {}
    for header, value in row.items():
        key = _map_header(dataset, header, meta.get("period_month") or "")
        if key and str(value or "").strip() and key not in mapped:
            mapped[key] = str(value).strip()
    period = mapped.get("period") or meta.get("period_month") or ""
    if mapped.get("trend_url") and not mapped.get("url"):
        mapped["url"] = mapped["trend_url"]
    if mapped.get("url") and not mapped.get("domain"):
        parsed = urlparse(mapped["url"] if "://" in mapped["url"] else f"https://{mapped['url']}")
        mapped["domain"] = (parsed.hostname or "").removeprefix("www.")
    mapped.update(
        {
            "period_month": period[:7] if len(period) >= 7 and period[4:5] == "-" else (meta.get("period_month") or ""),
            "previous_month": meta.get("previous_month") or "",
            "original_file_name": meta.get("original_file_name") or "",
            "relative_path": meta.get("relative_path") or "",
            "dataset_type": dataset,
            "provider": mapped.get("payment_provider") or meta.get("provider") or "",
            "source_note": mapped.get("note") or meta.get("source_note") or "",
        }
    )
    return mapped


def _clean_name(value: str) -> str:
    text = (value or "").replace("\\", "/").split("/")[-1].strip()
    return text or "upload"


def _clean_path(value: str, fallback: str) -> str:
    text = (value or "").replace("\\", "/").strip()
    parts = [part for part in text.split("/") if part and part not in {".", ".."}]
    return "/".join(parts) or fallback


def build_preview(files: list[tuple[str, str, bytes]], import_name: str, import_note: str, source_id: int | None, dataset_type: str, auto_detect: bool) -> dict:
    if not files:
        raise ValueError("请选择文件")
    if len(files) > _MAX_FILES:
        raise ValueError(f"一次最多 {_MAX_FILES} 个文件")
    override = (dataset_type or "").strip()
    if override and override not in DATASET_TYPES:
        raise ValueError("dataset_type 不在允许列表中")
    stored_files = []
    public_files = []
    for file_name, relative_path, data in files:
        name = _clean_name(file_name)
        path = _clean_path(relative_path, name)
        warnings = []
        headers: list[str] = []
        rows: list[dict] = []
        if len(data) > _MAX_BYTES:
            warnings.append("文件超过 8MB，未解析")
        elif not data:
            warnings.append("文件为空")
        else:
            try:
                headers, rows, extra = _read_table(name, data)
                warnings.extend(extra)
            except ValueError as exc:
                warnings.append(str(exc))
            except ImportError:
                warnings.append("XLSX 解析不可用")
        detected = detect_dataset(name, path, headers)
        if override and not auto_detect:
            record_type, provider = DATASET_TYPES[override]
            detected = {
                "detected_dataset_type": override,
                "record_type": record_type,
                "provider": provider,
                "confidence": 1.0,
                "warnings": warnings,
            }
        else:
            detected["warnings"] = warnings + detected["warnings"]
        if detected["confidence"] < 0.8:
            detected["warnings"].append("置信度低于 0.8，请确认后再导入")
        period, previous = _resolve_months(name, path, headers)
        if not period:
            detected["warnings"].append("未能识别月份")
        meta = {
            "period_month": period,
            "previous_month": previous,
            "original_file_name": name,
            "relative_path": path,
            "provider": detected["provider"],
            "source_note": (import_note or "").strip(),
        }
        dataset = detected["detected_dataset_type"]
        samples = [_map_row(dataset, row, meta) for row in rows[:3]] if dataset in DATASET_TYPES else rows[:3]
        stored_files.append(
            {
                "file_name": name,
                "relative_path": path,
                "headers": headers,
                "rows": rows,
                "detected_dataset_type": dataset,
                "record_type": detected["record_type"],
                "provider": detected["provider"],
                "confidence": detected["confidence"],
                "period_month": period,
                "previous_month": previous,
            }
        )
        public_files.append(
            {
                "file_name": name,
                "relative_path": path,
                "detected_dataset_type": dataset,
                "record_type": detected["record_type"],
                "provider": detected["provider"],
                "period_month": period,
                "previous_month": previous,
                "columns": headers,
                "sample_rows": samples,
                "confidence": detected["confidence"],
                "warnings": detected["warnings"],
            }
        )
    preview_id = uuid.uuid4().hex
    if len(_previews) >= 8:
        oldest = sorted(_previews, key=lambda key: _previews[key].get("created_at") or "")[0]
        _previews.pop(oldest, None)
    _previews[preview_id] = {
        "created_at": _now(),
        "import_name": (import_name or "").strip(),
        "import_note": (import_note or "").strip(),
        "source_id": source_id,
        "files": stored_files,
    }
    _stage_batch(preview_id, (import_name or "").strip(), (import_note or "").strip(), len(stored_files))
    return {"preview_id": preview_id, "status": "preview", "files": public_files}


def confirm_batch(payload: dict) -> dict:
    preview_id = (payload.get("preview_id") or "").strip()
    preview = _previews.get(preview_id)
    if preview is None:
        raise ValueError("预览已过期，请重新上传")
    low_confidence = any(float(item.get("confidence") or 0) < 0.8 for item in preview["files"])
    if low_confidence and not payload.get("review_confirmed"):
        raise ValueError("置信度低于 0.8，请确认后再导入")
    overrides = {}
    for item in payload.get("files") or []:
        key = ((item.get("relative_path") or "").strip(), (item.get("file_name") or "").strip())
        overrides[key] = item
    import_name = (payload.get("import_name") or preview.get("import_name") or "").strip()
    import_note = (payload.get("import_note") if payload.get("import_note") is not None else preview.get("import_note") or "")
    import_note = (import_note or "").strip()
    source_id = payload.get("source_id")
    if source_id in ("", None):
        source_id = preview.get("source_id")
    results = []
    for stored in preview["files"]:
        chosen = overrides.get((stored["relative_path"], stored["file_name"])) or {}
        dataset = (chosen.get("dataset_type") or "").strip() or stored["detected_dataset_type"]
        if dataset not in DATASET_TYPES:
            results.append(
                {
                    "file_name": stored["file_name"],
                    "import_id": 0,
                    "record_type": "",
                    "dataset_type": dataset,
                    "row_count": 0,
                    "status": "failed",
                    "message": "未能识别数据类型",
                }
            )
            continue
        default_type, default_provider = DATASET_TYPES[dataset]
        record_type = (chosen.get("record_type") or "").strip() or default_type
        if record_type not in _RECORD_TYPES:
            results.append(
                {
                    "file_name": stored["file_name"],
                    "import_id": 0,
                    "record_type": record_type,
                    "dataset_type": dataset,
                    "row_count": 0,
                    "status": "failed",
                    "message": "record_type 不在允许列表中",
                }
            )
            continue
        provider = chosen.get("provider")
        if provider is None or str(provider).strip() == "":
            provider = default_provider
        provider = str(provider).strip()[:80]
        file_note = (chosen.get("import_note") or "").strip() or import_note
        meta = {
            "period_month": stored["period_month"],
            "previous_month": stored["previous_month"],
            "original_file_name": stored["file_name"],
            "relative_path": stored["relative_path"],
            "provider": provider,
            "source_note": file_note,
        }
        mapped_rows = [_map_row(dataset, row, meta) for row in stored["rows"]]
        name = import_name or stored["file_name"]
        if import_name and len(preview["files"]) > 1:
            name = f"{import_name} / {stored['file_name']}"
        try:
            saved = _insert_import(
                dataset=dataset,
                record_type=record_type,
                provider=provider,
                source_id=int(source_id) if source_id else None,
                import_name=name,
                original_filename=stored["relative_path"] or stored["file_name"],
                notes=file_note,
                rows=mapped_rows,
            )
        except ValueError as exc:
            results.append(
                {
                    "file_name": stored["file_name"],
                    "import_id": 0,
                    "record_type": record_type,
                    "dataset_type": dataset,
                    "row_count": 0,
                    "status": "failed",
                    "message": str(exc),
                }
            )
            continue
        results.append(
            {
                "file_name": stored["file_name"],
                "import_id": saved["import_id"],
                "record_type": record_type,
                "dataset_type": dataset,
                "row_count": saved["row_count"],
                "status": saved["status"],
                "message": "",
            }
        )
    imported = any(item["status"] == "imported" for item in results)
    _finish_batch(preview_id, "imported" if imported else "failed")
    _previews.pop(preview_id, None)
    return {"imports": results}


def _stage_batch(preview_id: str, import_name: str, notes: str, file_count: int) -> None:
    conn = connect()
    try:
        conn.execute(
            """
            INSERT INTO import_batches (
                preview_id, import_name, status, file_count, notes, created_at, confirmed_at
            ) VALUES (?, ?, 'preview', ?, ?, ?, '')
            """,
            (preview_id, import_name, file_count, notes, _now()),
        )
        conn.commit()
    finally:
        conn.close()


def _finish_batch(preview_id: str, status: str) -> None:
    conn = connect()
    try:
        conn.execute(
            """
            UPDATE import_batches
            SET status = ?, confirmed_at = ?
            WHERE preview_id = ?
            """,
            (status, _now(), preview_id),
        )
        conn.commit()
    finally:
        conn.close()


def _named_source(name: str) -> dict | None:
    conn = connect()
    try:
        row = conn.execute("SELECT id FROM data_sources WHERE name = ?", (name,)).fetchone()
    finally:
        conn.close()
    if row is None:
        return None
    return get_source(int(row["id"]))


def _ensure_dataset_source(dataset: str) -> dict:
    name, source_type, provider = _DATASET_SOURCE[dataset]
    found = _named_source(name)
    if found:
        return found
    return create_source(
        {
            "name": name,
            "source_type": source_type,
            "provider": provider,
            "notes": f"dataset_type={dataset}",
            "enabled": 1,
        }
    )


def _insert_import(dataset: str, record_type: str, provider: str, source_id: int | None, import_name: str, original_filename: str, notes: str, rows: list[dict]) -> dict:
    if source_id:
        source = get_source(int(source_id))
        if source is None:
            raise ValueError("数据源不存在")
    else:
        source = _ensure_dataset_source(dataset)
    status = "imported" if rows else "empty"
    period = rows[0].get("period_month") if rows else ""
    note = f"dataset_type={dataset}; access_mode=upload; period_month={period}; {notes}".strip()
    now = _now()
    conn = connect()
    try:
        cur = conn.execute(
            """
            INSERT INTO source_imports (
                source_id, import_name, source_type, record_type, original_filename,
                row_count, status, notes, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(source["id"]),
                import_name or original_filename or "batch import",
                source.get("source_type") or "",
                record_type,
                original_filename,
                len(rows),
                status,
                note,
                now,
            ),
        )
        import_id = int(cur.lastrowid)
        for row in rows:
            domain = row.get("domain") or ""
            url = row.get("url") or ""
            title = row.get("title") or ""
            keyword = row.get("keyword") or ""
            metric_name = ""
            metric_value = ""
            for key in ("payment_traffic", "monthly_traffic", "current_traffic", "traffic", "current_dr", "search_volume", "score", "rank"):
                if row.get(key):
                    metric_name = key
                    metric_value = row[key]
                    break
            conn.execute(
                """
                INSERT INTO raw_source_records (
                    import_id, source_id, record_type, raw_json, normalized_title,
                    normalized_url, normalized_keyword, normalized_domain, metric_name,
                    metric_value, time_range, confidence, status, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    import_id,
                    int(source["id"]),
                    record_type,
                    json.dumps(row, ensure_ascii=False),
                    title,
                    url,
                    keyword,
                    domain,
                    metric_name,
                    metric_value,
                    row.get("period_month") or "",
                    "imported",
                    status,
                    now,
                ),
            )
        conn.commit()
    finally:
        conn.close()
    return {"import_id": import_id, "row_count": len(rows), "status": status}


def list_trend_signals(limit: int = 10) -> dict:
    size = max(1, min(int(limit or 10), 50))
    conn = connect()
    try:
        total = conn.execute(
            """
            SELECT COUNT(*) AS n
            FROM raw_source_records
            WHERE record_type = 'trend_signal'
              AND lower(COALESCE(status, '')) IN ('imported', 'confirmed')
            """
        ).fetchone()
        rows = conn.execute(
            """
            SELECT id, normalized_keyword, normalized_domain, normalized_url, time_range, raw_json, created_at
            FROM raw_source_records
            WHERE record_type = 'trend_signal'
              AND lower(COALESCE(status, '')) IN ('imported', 'confirmed')
            ORDER BY id DESC
            LIMIT ?
            """,
            (size,),
        ).fetchall()
    finally:
        conn.close()
    items = []
    for row in rows:
        try:
            raw = json.loads(row["raw_json"] or "{}")
        except json.JSONDecodeError:
            raw = {}
        if not isinstance(raw, dict):
            raw = {}
        items.append(
            {
                "id": int(row["id"]),
                "keyword": row["normalized_keyword"] or raw.get("keyword") or "",
                "search_volume": raw.get("search_volume") or "",
                "started_at": raw.get("started_at") or "",
                "growth": raw.get("growth_rate") or "",
                "related_queries": raw.get("related_queries") or "",
                "region": raw.get("region") or "",
                "period": row["time_range"] or raw.get("period_month") or raw.get("period") or "",
                "trend_url": raw.get("trend_url") or row["normalized_url"] or "",
                "domain": row["normalized_domain"] or "",
            }
        )
    return {"total": int(total["n"] or 0) if total else 0, "items": items}


def ignore_trend_signal(record_id: int) -> dict:
    conn = connect()
    try:
        row = conn.execute(
            "SELECT id FROM raw_source_records WHERE id = ? AND record_type = 'trend_signal'",
            (int(record_id),),
        ).fetchone()
        if row is None:
            raise ValueError("热词不存在")
        conn.execute(
            "UPDATE raw_source_records SET status = 'ignored' WHERE id = ?",
            (int(record_id),),
        )
        conn.commit()
    finally:
        conn.close()
    return {"id": int(record_id), "status": "ignored"}


def _sample_domain(domain: str) -> bool:
    host = (domain or "").lower().removeprefix("www.")
    return host in {"example.com", "test.com"} or host.endswith(".example.com") or host.endswith(".test.com")


def _crawl_summary() -> dict:
    conn = connect()
    try:
        rows = conn.execute(
            """
            SELECT normalized_title, normalized_url, normalized_domain, status, created_at
            FROM raw_source_records
            WHERE record_type = 'crawl_signal'
            ORDER BY id DESC
            LIMIT 8
            """
        ).fetchall()
        counts = conn.execute(
            """
            SELECT normalized_domain
            FROM raw_source_records
            WHERE record_type = 'crawl_signal'
            """
        ).fetchall()
    finally:
        conn.close()
    real = 0
    sample = 0
    for row in counts:
        if _sample_domain(row["normalized_domain"] or ""):
            sample += 1
        else:
            real += 1
    if real >= 10:
        status = "ready"
    elif real >= 1:
        status = "partial"
    else:
        status = "missing"
    return {
        "status": status,
        "real_count": real,
        "sample_count": sample,
        "recent": [dict(row) for row in rows],
    }


def _import_summary() -> dict:
    conn = connect()
    try:
        rows = conn.execute(
            """
            SELECT id, import_name, record_type, original_filename, row_count, status, created_at
            FROM source_imports
            ORDER BY id DESC
            LIMIT 8
            """
        ).fetchall()
        total = conn.execute("SELECT COUNT(*) AS n FROM source_imports").fetchone()
    finally:
        conn.close()
    count = int(total["n"] or 0)
    return {
        "status": "imported" if count else "missing",
        "batch_count": count,
        "recent": [dict(row) for row in rows],
    }


def _ledger() -> list[dict]:
    conn = connect()
    try:
        sources = conn.execute(
            """
            SELECT id, name, source_type, provider, notes, enabled
            FROM data_sources
            ORDER BY id ASC
            """
        ).fetchall()
        counts = conn.execute(
            """
            SELECT source_id, COUNT(*) AS row_count, MAX(created_at) AS last_run_at
            FROM raw_source_records
            GROUP BY source_id
            """
        ).fetchall()
        months = conn.execute(
            """
            SELECT source_id, time_range
            FROM raw_source_records
            WHERE time_range GLOB '[0-9][0-9][0-9][0-9]-[0-9][0-9]'
            """
        ).fetchall()
        imports = conn.execute(
            """
            SELECT source_id, status, notes, created_at, record_type
            FROM source_imports
            ORDER BY id DESC
            """
        ).fetchall()
    finally:
        conn.close()
    count_map = {int(row["source_id"]): row for row in counts}
    month_map: dict[int, set[str]] = {}
    for row in months:
        month_map.setdefault(int(row["source_id"]), set()).add(row["time_range"])
    latest_import = {}
    for row in imports:
        source_id = int(row["source_id"])
        if source_id not in latest_import:
            latest_import[source_id] = row
    health = {item["name"]: item for item in provider_health()["providers"]}
    ledger = []
    for source in sources:
        source_id = int(source["id"])
        source_type = source["source_type"] or ""
        if source_type in {"public_web", "news"}:
            access_mode = "crawl"
        elif source_type in {"serp", "search_console", "google_trends", "ai_search"}:
            access_mode = "api"
        else:
            access_mode = "upload"
        counted = count_map.get(source_id)
        latest = latest_import.get(source_id)
        provider_name = source["provider"] or ""
        dataset_type = ""
        if latest and latest["notes"]:
            match = re.search(r"dataset_type=([a-z0-9_]+)", latest["notes"])
            if match:
                dataset_type = match.group(1)
        if not dataset_type:
            if source_type == "payment_ranking" and "stripe" in f"{provider_name} {source['name'] or ''}".lower():
                dataset_type = "stripe_payment_ranking"
            elif source_type == "payment_ranking":
                dataset_type = "payment_ranking"
            else:
                dataset_type = {
                    "dr_growth": "dr_growth_ranking",
                    "site_traffic": "traffic_growth_ranking",
                    "new_site_growth": "new_website_ranking",
                    "public_web": "public_web_crawl",
                    "search_console": "google_search_console",
                    "serp": "serp_result",
                }.get(source_type, "")
        configured = bool(source["enabled"])
        if provider_name == "Serper":
            configured = bool(health.get("Serper", {}).get("configured"))
        if source_type == "search_console":
            configured = bool(health.get("GSC", {}).get("configured"))
        covered = sorted(month_map.get(source_id, set()))
        row_count = int(counted["row_count"]) if counted else 0
        last_status = latest["status"] if latest else ("ready" if row_count else "missing")
        ledger.append(
            {
                "source_name": source["name"] or "",
                "source_type": source_type,
                "provider": provider_name,
                "dataset_type": dataset_type,
                "access_mode": access_mode,
                "enabled": bool(source["enabled"]),
                "configured": configured,
                "last_run_at": (counted["last_run_at"] if counted and counted["last_run_at"] else "") or "",
                "last_status": last_status or "missing",
                "row_count": row_count,
                "months_covered": covered,
                "notes": source["notes"] or "",
            }
        )
    names = {row["source_name"] for row in ledger}
    if "GA4" not in names and not any(row["provider"] == "GA4" for row in ledger):
        ga4 = health.get("GA4", {})
        ledger.append(_planned_ledger("GA4", "google_analytics_4", "api", ga4))
    if not any(row["provider"] == "DataForSEO" or row["source_name"] == "DataForSEO" for row in ledger):
        ledger.append(_planned_ledger("DataForSEO", "dataforseo", "api", health.get("DataForSEO", {})))
    return ledger


def _planned_ledger(name: str, dataset_type: str, access_mode: str, health: dict) -> dict:
    return {
        "source_name": name,
        "source_type": dataset_type,
        "provider": name,
        "dataset_type": dataset_type,
        "access_mode": access_mode,
        "enabled": bool(health.get("enabled")),
        "configured": bool(health.get("configured")),
        "last_run_at": "",
        "last_status": health.get("status") or "not_configured",
        "row_count": 0,
        "months_covered": [],
        "notes": health.get("message") or "planned connector",
    }


def intake_overview() -> dict:
    health = provider_health()
    by_name = {item["name"]: item for item in health["providers"]}
    google = by_name.get("Google CSE", {})
    google_status = google.get("status") or "not_configured"
    serper = by_name.get("Serper", {})
    connectors = [
        {
            "provider": "serper",
            "name": "Serper SERP",
            "status": serper.get("status") or "not_configured",
            "configured": bool(serper.get("configured")),
            "message": serper.get("message") or "",
            "details": serper.get("details") or [],
            "uses": ["SERP Top 10", "Competitor Drafts"],
        },
        {
            "provider": "google_cse",
            "name": "Google CSE",
            "status": google_status,
            "configured": bool(google.get("configured")),
            "message": google.get("message") or "",
            "fallback": "serper / manual_csv",
            "details": google.get("details") or [],
            "uses": ["SERP Top 10"],
            "card": google_cse_card(),
        },
        _connector_from("gsc", "Google Search Console", by_name.get("GSC", {}), ["impressions", "clicks", "ctr", "position"]),
        _connector_from("ga4", "GA4", by_name.get("GA4", {}), ["sessions", "activeUsers", "conversions"]),
        _connector_from("dataforseo", "DataForSEO", by_name.get("DataForSEO", {}), ["SERP 批量", "keyword volume"]),
        _connector_from("sitedata", "SiteData CLI", by_name.get("SiteData", {}), ["traffic_growth", "domain_rating_growth", "payment_traffic"]),
    ]
    crawl = _crawl_summary()
    imports = _import_summary()
    ledger = _ledger()
    return {
        "crawl": crawl,
        "connectors": connectors,
        "imports": imports,
        "ledger": ledger,
        "serp": serp_choices(),
    }


def _connector_from(provider: str, name: str, row: dict, uses: list[str]) -> dict:
    return {
        "provider": provider,
        "name": name,
        "status": row.get("status") or "not_configured",
        "configured": bool(row.get("configured")),
        "enabled": bool(row.get("enabled")),
        "message": row.get("message") or "planned connector",
        "details": row.get("details") or [],
        "uses": uses,
    }


def gsc_health() -> dict:
    row = next(item for item in provider_health()["providers"] if item["name"] == "GSC")
    return {
        "provider": "google_search_console",
        "enabled": row["enabled"],
        "configured": row["configured"],
        "status": row["status"],
        "message": row["message"],
        "auth_mode": _detail_value(row, "auth_mode"),
        "site_url": _detail_value(row, "site_url"),
        "record_type": "validation_signal",
        "dataset_type": "google_search_console",
    }


def ga4_health() -> dict:
    row = next(item for item in provider_health()["providers"] if item["name"] == "GA4")
    return {
        "provider": "google_analytics_4",
        "enabled": row["enabled"],
        "configured": row["configured"],
        "status": row["status"],
        "message": row["message"],
        "auth_mode": _detail_value(row, "auth_mode"),
        "property_id": _detail_value(row, "property_id"),
        "record_type": "validation_signal",
        "dataset_type": "google_analytics_4",
    }


def dataforseo_health() -> dict:
    row = next(item for item in provider_health()["providers"] if item["name"] == "DataForSEO")
    return {
        "provider": "dataforseo",
        "enabled": row["enabled"],
        "configured": row["configured"],
        "status": row["status"],
        "message": row["message"],
        "login_present": _detail_value(row, "login_present") == "true",
        "password_present": _detail_value(row, "password_present") == "true",
        "location_code": _detail_value(row, "location_code"),
        "language_code": _detail_value(row, "language_code"),
    }


def _detail_value(row: dict, key: str) -> str:
    prefix = f"{key}="
    for item in row.get("details") or []:
        if str(item).startswith(prefix):
            return str(item)[len(prefix):]
    return ""


def _public_url(value: str) -> str:
    parsed = urlparse((value or "").strip())
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("只抓取公开 http(s) 页面")
    host = parsed.hostname.lower()
    if host in {"localhost", "127.0.0.1", "::1"}:
        raise ValueError("不抓取本机地址")
    try:
        ip = ipaddress.ip_address(host)
    except ValueError:
        ip = None
    if ip is not None and (ip.is_private or ip.is_loopback or ip.is_link_local):
        raise ValueError("不抓取内网地址")
    parts = [part for part in (parsed.path or "").lower().split("/") if part]
    if any(part in {"login", "signin", "sign-in", "wp-admin", "wp-login", "admin"} for part in parts):
        raise ValueError("不抓取登录或后台页面")
    return parsed.geturl()


def _page_text(html: str) -> str:
    cleaned = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", html or "")
    cleaned = re.sub(r"(?s)<[^>]+>", " ", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned[:800]


def _tag_text(html: str, pattern: str) -> str:
    match = re.search(pattern, html or "", re.I | re.S)
    if not match:
        return ""
    return re.sub(r"\s+", " ", match.group(1)).strip()


def _parse_page(url: str, status_code: int, html: str, keyword: str, crawl_type: str) -> dict:
    title = _tag_text(html, r"<title[^>]*>(.*?)</title>")
    description = ""
    desc = re.search(r'<meta[^>]+name=["\']description["\'][^>]+content=["\'](.*?)["\']', html or "", re.I | re.S)
    if desc is None:
        desc = re.search(r'<meta[^>]+content=["\'](.*?)["\'][^>]+name=["\']description["\']', html or "", re.I | re.S)
    if desc:
        description = re.sub(r"\s+", " ", desc.group(1)).strip()
    h1 = _tag_text(html, r"<h1[^>]*>(.*?)</h1>")
    canonical = ""
    canon = re.search(r'<link[^>]+rel=["\']canonical["\'][^>]+href=["\'](.*?)["\']', html or "", re.I)
    if canon:
        canonical = canon.group(1).strip()
    links = []
    for href in re.findall(r'href=["\'](https?://[^"\']+)["\']', html or "", re.I):
        if href not in links and href.rstrip("/") != url.rstrip("/"):
            links.append(href)
        if len(links) >= 30:
            break
    text = _page_text(html)
    lowered = f"{html} {text}".lower()
    providers = [word for word in _PAYMENT_WORDS if word in lowered]
    now = _now()
    return {
        "url": url,
        "status_code": status_code,
        "title": title,
        "meta_description": description,
        "h1": h1,
        "canonical": canonical,
        "outbound_links": links,
        "detected_cta": any(token in lowered for token in ("sign up", "signup", "get started", "start free", "try free")),
        "detected_pricing": any(token in lowered for token in ("pricing", "/mo", "per month")),
        "detected_signup": any(token in lowered for token in ("sign up", "signup", "register", "create account")),
        "detected_payment_provider": ", ".join(providers),
        "page_text_sample": text,
        "crawled_at": now,
        "keyword": keyword,
        "crawl_type": crawl_type,
        "dataset_type": "public_web_crawl",
        "record_type": "crawl_signal",
    }


def _fetch(url: str) -> tuple[int, str, str]:
    started = time.perf_counter()
    status = 0
    final_url = url
    try:
        with httpx.Client(timeout=15, follow_redirects=True, max_redirects=5) as client:
            response = client.get(url, headers={"User-Agent": "PJIntelligence/0.6 (public-page crawl)"})
        final_url = _public_url(str(response.url))
        status = response.status_code
        body = response.text[:1_000_000]
    except ValueError:
        raise
    except httpx.HTTPError as exc:
        raise ValueError("公开页面抓取失败") from exc
    finally:
        elapsed = int((time.perf_counter() - started) * 1000)
        _log.info("provider=crawl url=%s status_code=%s elapsed_ms=%s", url, status, elapsed)
        record_trace("GET", "external_crawl", status or 502, elapsed, "external_crawl", f"provider=crawl status_code={status or 502}")
    return status, final_url, body


def _sitemap_urls(xml: str, page_url: str) -> list[str]:
    host = urlparse(page_url).hostname
    found = []
    for loc in re.findall(r"<loc>\s*(https?://[^<\s]+)\s*</loc>", xml or "", re.I):
        try:
            clean = _public_url(loc.strip())
        except ValueError:
            continue
        if urlparse(clean).hostname != host:
            continue
        if clean not in found:
            found.append(clean)
        if len(found) >= _SITEMAP_CAP:
            break
    return found


def run_crawl(url: str, domain: str, keyword: str, crawl_type: str) -> dict:
    kind = (crawl_type or "").strip()
    if kind not in CRAWL_TYPES:
        raise ValueError("crawl_type 不在允许列表中")
    target = (url or "").strip()
    if not target and (domain or "").strip():
        host = domain.strip()
        host = host if "://" in host else f"https://{host}"
        target = host
    if not target:
        raise ValueError("公开抓取需要 url 或 domain")
    target = _public_url(target)
    status, final_url, body = _fetch(target)
    pages = [(status, final_url, body)]
    truncated = False
    if kind == "sitemap":
        extra = _sitemap_urls(body, final_url)
        if len(re.findall(r"<loc>", body or "", re.I)) > _SITEMAP_CAP:
            truncated = True
        for loc in extra:
            if loc.rstrip("/") == final_url.rstrip("/"):
                continue
            try:
                pages.append(_fetch(loc))
            except ValueError:
                continue
            if len(pages) >= _SITEMAP_CAP:
                truncated = True
                break
    source = _ensure_dataset_source("public_web_crawl")
    now = _now()
    records = [_parse_page(page_url, code, html, (keyword or "").strip(), kind) for code, page_url, html in pages]
    job_status = "partial" if truncated else "imported"
    if records and all(int(row["status_code"] or 0) >= 400 for row in records):
        job_status = "failed"
    conn = connect()
    try:
        cur = conn.execute(
            """
            INSERT INTO source_imports (
                source_id, import_name, source_type, record_type, original_filename,
                row_count, status, notes, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(source["id"]),
                f"Crawl {kind}",
                source.get("source_type") or "public_web",
                "crawl_signal",
                final_url,
                len(records),
                job_status,
                f"dataset_type=public_web_crawl; access_mode=crawl; keyword={keyword or ''}",
                now,
            ),
        )
        import_id = int(cur.lastrowid)
        for row in records:
            host = urlparse(row["url"]).hostname or ""
            if host.startswith("www."):
                host = host[4:]
            conn.execute(
                """
                INSERT INTO raw_source_records (
                    import_id, source_id, record_type, raw_json, normalized_title,
                    normalized_url, normalized_keyword, normalized_domain, metric_name,
                    metric_value, time_range, confidence, status, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    import_id,
                    int(source["id"]),
                    "crawl_signal",
                    json.dumps(row, ensure_ascii=False),
                    row.get("title") or "",
                    row.get("url") or "",
                    keyword or "",
                    host,
                    "status_code",
                    str(row.get("status_code") or ""),
                    row.get("crawled_at") or "",
                    "crawled",
                    "imported" if int(row.get("status_code") or 0) < 400 else "failed",
                    now,
                ),
            )
        conn.commit()
    finally:
        conn.close()
    public_records = []
    for row in records:
        shown = dict(row)
        shown["page_text_sample"] = (shown.get("page_text_sample") or "")[:180]
        shown["outbound_links"] = (shown.get("outbound_links") or [])[:8]
        public_records.append(shown)
    return {
        "status": job_status,
        "crawl_type": kind,
        "import_id": import_id,
        "pages": len(records),
        "records": public_records,
    }
